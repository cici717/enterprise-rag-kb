import hashlib
import json
import re
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator, Sequence

import numpy as np
from openai import OpenAI
from sqlalchemy import select
from sqlalchemy.orm import Session

from .config import Settings, get_settings
from .db import Chunk, Conversation, Document, Message, SessionLocal

TOKEN_RE = re.compile(r"[A-Za-z0-9_]+|[\u4e00-\u9fff]")


def tokenize(text: str) -> list[str]:
    base = TOKEN_RE.findall((text or "").lower())
    tokens = list(base)
    chinese_chars = [t for t in base if len(t) == 1 and "\u4e00" <= t <= "\u9fff"]
    tokens.extend(a + b for a, b in zip(chinese_chars, chinese_chars[1:]))
    return tokens


class HashingEmbeddingProvider:
    name = "hashing"

    def __init__(self, dim: int = 384) -> None:
        self.dim = dim

    def _embed_one(self, text: str) -> list[float]:
        vec = np.zeros(self.dim, dtype=np.float32)
        for token in tokenize(text):
            digest = hashlib.md5(token.encode("utf-8")).hexdigest()
            idx = int(digest[:8], 16) % self.dim
            sign = 1.0 if int(digest[8:10], 16) % 2 == 0 else -1.0
            vec[idx] += sign
        norm = float(np.linalg.norm(vec))
        if norm > 0:
            vec /= norm
        return vec.tolist()

    def embed(self, texts: Sequence[str]) -> list[list[float]]:
        return [self._embed_one(text) for text in texts]


class OpenAIEmbeddingProvider:
    name = "openai"

    def __init__(
        self,
        api_key: str,
        base_url: str,
        model: str,
        dim: int | None = None,
    ) -> None:
        self.client = OpenAI(api_key=api_key, base_url=base_url or None)
        self.model = model
        self.dim = dim

    def embed(self, texts: Sequence[str]) -> list[list[float]]:
        response = self.client.embeddings.create(
            model=self.model,
            input=list(texts),
        )
        return [item.embedding for item in response.data]


def get_embedding_provider(settings: Settings):
    provider = settings.embedding_provider.lower()
    api_key = settings.embedding_api_key or settings.llm_api_key
    if provider == "openai" and api_key:
        return OpenAIEmbeddingProvider(
            api_key=api_key,
            base_url=settings.embedding_base_url or settings.llm_base_url,
            model=settings.embedding_model,
            dim=settings.embedding_dim,
        )
    return HashingEmbeddingProvider(settings.embedding_dim)


def cosine_similarity(a: Sequence[float], b: Sequence[float]) -> float:
    va = np.asarray(a, dtype=np.float32)
    vb = np.asarray(b, dtype=np.float32)
    if va.size == 0 or vb.size == 0 or va.shape != vb.shape:
        return 0.0
    denom = float(np.linalg.norm(va) * np.linalg.norm(vb))
    if denom == 0:
        return 0.0
    return float(np.dot(va, vb) / denom)


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 120) -> list[str]:
    text = re.sub(r"\r\n?", "\n", text or "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    units: list[str] = []

    for paragraph in paragraphs:
        if len(paragraph) <= chunk_size:
            units.append(paragraph)
            continue
        sentences = re.split(r"(?<=[。！？!?；;])\s*", paragraph)
        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence:
                continue
            if len(sentence) <= chunk_size:
                units.append(sentence)
            else:
                for start in range(0, len(sentence), chunk_size):
                    units.append(sentence[start : start + chunk_size])

    chunks: list[str] = []
    current = ""
    for unit in units:
        candidate = f"{current}\n{unit}".strip() if current else unit
        if len(candidate) <= chunk_size:
            current = candidate
        else:
            if current:
                chunks.append(current)
            current = unit
    if current:
        chunks.append(current)

    if overlap > 0 and len(chunks) > 1:
        result = [chunks[0]]
        for chunk in chunks[1:]:
            tail = result[-1][-overlap:]
            result.append((tail + "\n" + chunk).strip())
        return result
    return chunks


def parse_document(path: Path) -> list[tuple[int | None, str]]:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [
            (index + 1, page.extract_text() or "")
            for index, page in enumerate(reader.pages)
        ]

    if suffix == ".docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [(None, text)]

    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return [(None, path.read_text(encoding=encoding))]
        except UnicodeDecodeError:
            continue
    return [(None, path.read_text(encoding="utf-8", errors="ignore"))]


def save_upload(data: bytes, filename: str, upload_dir: Path) -> Path:
    safe_name = Path(filename or "document.txt").name
    target = upload_dir / f"{uuid.uuid4().hex}_{safe_name}"
    target.write_bytes(data)
    return target


def ingest_document(
    db: Session,
    filename: str,
    content_type: str,
    data: bytes,
    settings: Settings | None = None,
) -> Document:
    settings = settings or get_settings()
    path = save_upload(data, filename, settings.upload_dir)

    document = Document(
        filename=Path(filename or "document.txt").name,
        content_type=content_type,
        size=len(data),
        status="processing",
    )
    db.add(document)
    db.commit()
    db.refresh(document)

    try:
        pages = parse_document(path)
        chunk_rows: list[tuple[int | None, str]] = []
        for page, text in pages:
            for part in chunk_text(text, settings.chunk_size, settings.chunk_overlap):
                if part.strip():
                    chunk_rows.append((page, part.strip()))

        if not chunk_rows:
            raise ValueError("文档未解析出可用文本")

        texts = [content for _, content in chunk_rows]
        embeddings = get_embedding_provider(settings).embed(texts)

        for index, ((page, content), embedding) in enumerate(
            zip(chunk_rows, embeddings)
        ):
            db.add(
                Chunk(
                    document_id=document.id,
                    chunk_index=index,
                    content=content,
                    page=page,
                    embedding=json.dumps(embedding, ensure_ascii=False),
                )
            )

        document.chunk_count = len(chunk_rows)
        document.status = "ready"
        db.commit()
        db.refresh(document)
        return document
    except Exception:
        document.status = "failed"
        db.commit()
        raise


@dataclass
class RetrievedChunk:
    chunk_id: int
    document_id: int
    filename: str
    page: int | None
    chunk_index: int
    content: str
    semantic_score: float
    lexical_score: float
    score: float


def retrieve(
    db: Session,
    query: str,
    top_k: int = 5,
    settings: Settings | None = None,
) -> list[RetrievedChunk]:
    settings = settings or get_settings()
    provider = get_embedding_provider(settings)
    query_vector = provider.embed([query])[0]
    query_tokens = set(tokenize(query))

    rows = db.execute(
        select(Chunk, Document.filename).join(
            Document,
            Chunk.document_id == Document.id,
        )
    ).all()

    results: list[RetrievedChunk] = []
    for chunk, filename in rows:
        try:
            embedding = json.loads(chunk.embedding or "[]")
        except json.JSONDecodeError:
            embedding = []

        semantic = cosine_similarity(query_vector, embedding)
        chunk_tokens = set(tokenize(chunk.content))
        lexical = (
            len(query_tokens & chunk_tokens) / max(1, len(query_tokens))
            if query_tokens
            else 0.0
        )
        score = 0.8 * semantic + 0.2 * lexical
        results.append(
            RetrievedChunk(
                chunk_id=chunk.id,
                document_id=chunk.document_id,
                filename=filename,
                page=chunk.page,
                chunk_index=chunk.chunk_index,
                content=chunk.content,
                semantic_score=semantic,
                lexical_score=lexical,
                score=score,
            )
        )

    results.sort(key=lambda item: item.score, reverse=True)
    return results[: max(1, top_k)]


def citations_payload(retrieved: list[RetrievedChunk]) -> list[dict]:
    return [
        {
            "document_id": item.document_id,
            "filename": item.filename,
            "page": item.page,
            "chunk_index": item.chunk_index,
            "score": round(item.score, 4),
            "snippet": item.content[:220],
        }
        for item in retrieved
    ]


def build_messages(query: str, retrieved: list[RetrievedChunk]) -> list[dict]:
    context_parts = []
    for index, item in enumerate(retrieved, start=1):
        page_text = f" 第{item.page}页" if item.page else ""
        context_parts.append(
            f"[{index}] 来源: {item.filename}{page_text}\n{item.content}"
        )
    context = "\n\n".join(context_parts) if context_parts else "（没有检索到相关资料）"

    system_prompt = (
        "你是一个企业知识库助手。只能根据提供的参考资料回答，"
        "不要编造资料中没有的信息。如果资料不足，请明确说不知道。"
        "回答要简洁、准确，并在相关句子后标注引用编号，例如 [1]。"
    )
    user_prompt = (
        f"参考资料：\n{context}\n\n"
        f"用户问题：{query}\n\n"
        "请用中文回答，并给出引用编号。"
    )
    return [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_prompt},
    ]


class OpenAICompatibleLLM:
    mode = "openai"

    def __init__(self, api_key: str, base_url: str, model: str) -> None:
        self.client = OpenAI(api_key=api_key, base_url=base_url or None)
        self.model = model

    def complete(self, messages: list[dict]) -> str:
        response = self.client.chat.completions.create(
            model=self.model,
            messages=messages,
            temperature=0.2,
        )
        return response.choices[0].message.content or ""

    def stream(self, messages: list[dict]) -> Iterator[str]:
        stream = self.client.chat.completions.create(
            model=self.model,
            messages=messages,
            temperature=0.2,
            stream=True,
        )
        for chunk in stream:
            if not chunk.choices:
                continue
            delta = chunk.choices[0].delta.content
            if delta:
                yield delta


class MockLLM:
    mode = "mock"

    def _fallback(self, messages: list[dict]) -> str:
        user_content = messages[-1]["content"] if messages else ""
        context = user_content.split("参考资料：", 1)[-1]
        context = context.split("用户问题：", 1)[0].strip()
        if len(context) > 1500:
            context = context[:1500] + "..."
        return (
            "当前未配置大模型 API，已使用离线模式回答。\n\n"
            "根据检索到的资料，相关内容如下：\n\n"
            f"{context or '暂未检索到相关资料。'}\n\n"
            "配置 .env 中的 LLM_API_KEY 后，可由大模型生成更自然的回答。"
        )

    def complete(self, messages: list[dict]) -> str:
        return self._fallback(messages)

    def stream(self, messages: list[dict]) -> Iterator[str]:
        text = self._fallback(messages)
        for start in range(0, len(text), 18):
            yield text[start : start + 18]


def get_llm(settings: Settings):
    if settings.llm_api_key:
        return OpenAICompatibleLLM(
            api_key=settings.llm_api_key,
            base_url=settings.llm_base_url,
            model=settings.llm_model,
        )
    return MockLLM()


def ensure_conversation(
    db: Session,
    conversation_id: int | None,
    query: str,
) -> Conversation:
    if conversation_id:
        existing = db.get(Conversation, conversation_id)
        if existing:
            return existing
    conversation = Conversation(title=(query[:30] or "新会话"))
    db.add(conversation)
    db.commit()
    db.refresh(conversation)
    return conversation


def chat(
    db: Session,
    query: str,
    conversation_id: int | None = None,
    top_k: int | None = None,
    settings: Settings | None = None,
) -> tuple[str, list[dict], int]:
    settings = settings or get_settings()
    retrieved = retrieve(db, query, top_k or settings.top_k, settings)
    messages = build_messages(query, retrieved)
    answer = get_llm(settings).complete(messages)
    citations = citations_payload(retrieved)

    conversation = ensure_conversation(db, conversation_id, query)
    db.add(Message(conversation_id=conversation.id, role="user", content=query))
    db.add(
        Message(
            conversation_id=conversation.id,
            role="assistant",
            content=answer,
            citations=json.dumps(citations, ensure_ascii=False),
        )
    )
    db.commit()
    return answer, citations, conversation.id


def _sse(payload: dict) -> str:
    return f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"


def stream_chat(
    query: str,
    conversation_id: int | None = None,
    top_k: int | None = None,
    settings: Settings | None = None,
) -> Iterator[str]:
    settings = settings or get_settings()
    db = SessionLocal()
    try:
        retrieved = retrieve(db, query, top_k or settings.top_k, settings)
        messages = build_messages(query, retrieved)
        citations = citations_payload(retrieved)
        conversation = ensure_conversation(db, conversation_id, query)

        yield _sse({"type": "conversation", "data": {"id": conversation.id}})
        yield _sse({"type": "citations", "data": citations})

        answer_parts: list[str] = []
        for token in get_llm(settings).stream(messages):
            answer_parts.append(token)
            yield _sse({"type": "token", "data": token})

        answer = "".join(answer_parts)
        db.add(Message(conversation_id=conversation.id, role="user", content=query))
        db.add(
            Message(
                conversation_id=conversation.id,
                role="assistant",
                content=answer,
                citations=json.dumps(citations, ensure_ascii=False),
            )
        )
        db.commit()
        yield _sse({"type": "done", "data": {}})
    except Exception as exc:
        db.rollback()
        yield _sse({"type": "error", "data": {"message": str(exc)}})
    finally:
        db.close()